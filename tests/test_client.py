from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import httpx
import pytest
import respx
from openpyxl import Workbook

from kric import (
    KricAuthError,
    KricClient,
    KricFileClient,
    KricInvalidParameterError,
    KricNetworkError,
    KricRateLimitError,
    KricServerError,
    ServiceDayCode,
    parse_nationwide_station_info_xlsx,
    parse_xlsx_table,
)
from kric.files import FILE_DOWNLOAD_URL
from kric.parse import (
    extract_items,
    float_or_none,
    integer_or_none,
    parse_station_facility,
    parse_station_info,
    parse_subway_timetable_entry,
    service_day_code,
)


@respx.mock
async def test_station_info_sends_documented_json_parameters():
    route = respx.get("https://openapi.kric.go.kr/openapi/convenientInfo/stationInfo").mock(
        return_value=httpx.Response(200, json={"header": {"resultCode": "00"}, "body": {"items": {"item": {
            "railOprIsttCd": "S1", "lnCd": "1", "stinCd": "150", "stinNm": "서울역",
            "stinLocLat": "37.554648", "stinLocLon": "126.970606", "mapCordX": "198000", "mapCordY": "451000",
        }}}})
    )
    async with KricClient("test-key") as client:
        rows = await client.get_station_info(rail_operator_code="S1", line_code="1", station_code="150")

    assert route.called
    assert route.calls[0].request.url.params["format"] == "json"
    assert route.calls[0].request.url.params["serviceKey"] == "test-key"
    assert rows[0].station_code == "150"
    assert rows[0].latitude == 37.554648
    assert rows[0].map_x == "198000"


@respx.mock
async def test_route_keeps_route_groups_order_and_repeated_station():
    respx.get("https://openapi.kric.go.kr/openapi/trainUseInfo/subwayRouteInfo").mock(
        return_value=httpx.Response(200, json={"header": {"resultCode": "00"}, "body": {"items": {"item": [
            {"mreaWideCd": "01", "railOprIsttCd": "S1", "lnCd": "2", "routCd": "R2", "stinCd": "202", "stinNm": "분기역", "stinConsOrdr": "1"},
            {"mreaWideCd": "01", "railOprIsttCd": "S1", "lnCd": "2", "routCd": "R1", "stinCd": "201", "stinNm": "순환역", "stinConsOrdr": "44"},
            {"mreaWideCd": "01", "railOprIsttCd": "S1", "lnCd": "2", "routCd": "R1", "stinCd": "201", "stinNm": "순환역", "stinConsOrdr": "1"},
        ]}}})
    )
    async with KricClient("test-key") as client:
        rows = await client.get_subway_route_info(metro_area_code="01", line_code="2")

    assert [(row.route_code, row.station_code, row.station_sequence) for row in rows] == [
        ("R2", "202", 1), ("R1", "201", 1), ("R1", "201", 44)
    ]


@respx.mock
async def test_timetable_preserves_time_strings_and_day_code():
    respx.get("https://openapi.kric.go.kr/openapi/convenientInfo/stationTimetable").mock(
        return_value=httpx.Response(200, json={"header": {"resultCode": "00"}, "body": {"items": {"item": {
            "railOprIsttCd": "S1", "lnCd": "3", "stinCd": "312", "dayCd": "8", "dayNm": "평일",
            "arvTm": "05:31", "dptTm": "05:32", "orgStinCd": "310", "tmnStinCd": "330", "trnNo": "1001",
        }}}})
    )
    async with KricClient("test-key") as client:
        rows = await client.get_station_timetable(
            rail_operator_code="S1", line_code="3", station_code="312", day_code=ServiceDayCode.WEEKDAY
        )

    row = rows[0]
    assert row.day_code is ServiceDayCode.WEEKDAY
    assert (row.arrival_time, row.departure_time, row.origin_station_code, row.terminal_station_code) == (
        "05:31", "05:32", "310", "330"
    )


@respx.mock
async def test_error_envelope_maps_invalid_service_key():
    respx.get("https://openapi.kric.go.kr/openapi/convenientInfo/stationInfo").mock(
        return_value=httpx.Response(200, json={"header": {"resultCode": "30", "resultMsg": "등록되지 않은 서비스키입니다."}})
    )
    async with KricClient("invalid") as client:
        with pytest.raises(KricAuthError, match="서비스키"):
            await client.get_station_info(station_code="150")


@respx.mock
async def test_error_envelope_maps_quota_message():
    respx.get("https://openapi.kric.go.kr/openapi/convenientInfo/stationInfo").mock(
        return_value=httpx.Response(200, json={"header": {"resultCode": "99", "resultMsg": "호출 횟수 초과"}})
    )
    async with KricClient("test-key") as client:
        with pytest.raises(KricRateLimitError):
            await client.get_station_info(station_code="150")


@respx.mock
@pytest.mark.parametrize(
    ("status_code", "error_type"),
    [(401, KricAuthError), (429, KricRateLimitError), (500, KricServerError)],
)
async def test_http_errors_are_not_treated_as_empty_results(status_code, error_type):
    respx.get("https://openapi.kric.go.kr/openapi/convenientInfo/stationInfo").mock(
        return_value=httpx.Response(status_code)
    )
    async with KricClient("test-key") as client:
        with pytest.raises(error_type):
            await client.get_station_info(station_code="150")


@respx.mock
async def test_non_json_and_network_errors_are_explicit():
    route = respx.get("https://openapi.kric.go.kr/openapi/convenientInfo/stationInfo")
    route.mock(return_value=httpx.Response(200, text="not-json"))
    async with KricClient("test-key") as client:
        with pytest.raises(KricServerError, match="not JSON"):
            await client.get_station_info(station_code="150")

    route.mock(side_effect=httpx.ConnectError("offline"))
    async with KricClient("test-key") as client:
        with pytest.raises(KricNetworkError):
            await client.get_station_info(station_code="150")


@respx.mock
async def test_subway_timetable_and_facilities_keep_documented_codes():
    timetable = respx.get("https://openapi.kric.go.kr/openapi/trainUseInfo/subwayTimetable").mock(
        return_value=httpx.Response(200, json={"header": {"resultCode": "00"}, "response": {"body": {"items": {"item": [{
            "railOprIsttCd": "S1", "lnCd": "1", "stinCd": "0150", "dayCd": "9", "arvTm": "24:01", "dptTm": "24:02", "trnNo": "0007"
        }]}}}})
    )
    facilities = respx.get("https://openapi.kric.go.kr/openapi/convenientInfo/stationCnvFacl").mock(
        return_value=httpx.Response(200, json={"header": {"resultCode": "00"}, "body": {"items": {"item": {
            "railOprIsttCd": "S1", "lnCd": "1", "stinCd": "0150", "facilityCode": "ELV", "unknown": "keep"
        }}}})
    )
    async with KricClient("test-key") as client:
        timetable_rows = await client.get_subway_timetable(
            rail_operator_code="S1", line_code="1", station_code="0150", day_code="9"
        )
        facility_rows = await client.get_station_facilities(
            rail_operator_code="S1", line_code="1", station_code="0150"
        )

    assert timetable.called and facilities.called
    assert timetable_rows[0].day_code is ServiceDayCode.HOLIDAY
    assert timetable_rows[0].departure_time == "24:02"
    assert timetable_rows[0].station_code == "0150"
    assert facility_rows[0].values["unknown"] == "keep"


def test_parser_rejects_bad_shapes_and_numeric_contracts():
    assert extract_items({"items": {"item": {"code": "001"}}}) == ({"code": "001"},)
    assert float_or_none(" ", "latitude") is None
    assert integer_or_none(None, "order") is None
    assert service_day_code("new-code") == "new-code"
    assert parse_subway_timetable_entry({
        "railOprIsttCd": "S1", "lnCd": "1", "stinCd": "0150", "dayCd": "7", "trnNo": "001"
    }).train_number == "001"
    assert parse_station_facility({
        "railOprIsttCd": "S1", "lnCd": "1", "stinCd": "0150", "blank": " "
    }).values["blank"] is None
    for invalid_coordinate in ("north", "NaN", "Infinity", "-Infinity"):
        with pytest.raises(KricServerError, match="numeric"):
            float_or_none(invalid_coordinate, "latitude")
    with pytest.raises(KricServerError, match="integer"):
        integer_or_none("first", "order")
    with pytest.raises(KricServerError, match="does not contain"):
        extract_items({"header": {"resultCode": "00"}})
    with pytest.raises(KricServerError, match="stationInfo item"):
        parse_station_info({"stinCd": "150", "stinNm": "서울역"})


async def test_client_rejects_blank_or_unsafe_requests_before_network():
    with pytest.raises(KricInvalidParameterError):
        KricClient(" ")
    async with KricClient("test-key") as client:
        with pytest.raises(KricInvalidParameterError):
            await client.get_station_info()
        with pytest.raises(KricInvalidParameterError):
            await client.get_subway_route_info(metro_area_code=" ", line_code="1")
        with pytest.raises(KricInvalidParameterError):
            await client.get_station_timetable(
                rail_operator_code="S1", line_code="1", station_code="150", day_code="0"
            )
    with pytest.raises(KricInvalidParameterError):
        KricClient("test-key", timeout=0)


def _station_info_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([
        "철도운영기관명", "운영노선", "역 종류", "역 번호", "역명(한글)", "역명(영어)",
        "역 위치(경도)", "역 위치(위도)", "역 주소(도로명 주소)", "데이터 기준일자", "추가 열",
    ])
    sheet.append([
        "서울교통공사", "1호선", "일반역", "0150", "서울역", "Seoul", 126.970606, 37.554648,
        "서울특별시 용산구", "2026-07-01", "원문 보존",
    ])
    sheet.cell(2, 4).value = 150
    sheet.cell(2, 4).number_format = "0000"
    sheet.append([
        "한국철도공사", "경의중앙선", "일반역", 110, "가상역", None, None, None, None, None, None,
    ])
    sheet.cell(3, 4).number_format = "\\I000"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_public_station_file_parser_keeps_file_identity_and_unknown_columns():
    rows = parse_nationwide_station_info_xlsx(_station_info_workbook_bytes())

    assert len(rows) == 2
    row = rows[0]
    assert row.rail_operator_name == "서울교통공사"
    assert row.operating_line_name == "1호선"
    assert row.station_number == "0150"
    assert row.station_name == "서울역"
    assert (row.longitude, row.latitude) == (126.970606, 37.554648)
    assert row.raw["추가 열"] == "원문 보존"
    assert rows[1].station_number == "I110"

    table = parse_xlsx_table(_station_info_workbook_bytes())
    assert table.headers[-1] == "추가 열"
    assert table.rows[0]["추가 열"] == "원문 보존"


@respx.mock
async def test_public_file_download_requires_no_service_key_and_parses_station_dataset():
    route = respx.get(FILE_DOWNLOAD_URL).mock(
        return_value=httpx.Response(
            200,
            content=_station_info_workbook_bytes(),
            headers={"content-type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
        )
    )
    async with KricFileClient() as client:
        rows = await client.get_nationwide_station_info()

    assert route.called
    params = route.calls[0].request.url.params
    assert dict(params) == {"type": "filedata", "id": "1294", "operation": "1"}
    assert "serviceKey" not in params
    assert rows[0].station_name == "서울역"


def test_public_station_file_parser_rejects_missing_contract_headers():
    workbook = Workbook()
    workbook.active.append(["운영노선", "역 번호", "역명(한글)"])
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(KricServerError, match="required headers"):
        parse_nationwide_station_info_xlsx(output.getvalue())


def test_public_file_parser_rejects_invalid_or_oversized_input():
    content = _station_info_workbook_bytes()

    with pytest.raises(KricInvalidParameterError, match="content must be bytes"):
        parse_xlsx_table("not-bytes")  # type: ignore[arg-type]
    with pytest.raises(KricServerError, match="max_compressed_bytes"):
        parse_xlsx_table(content, max_compressed_bytes=1)
    with pytest.raises(KricInvalidParameterError, match="HTTPS URL"):
        KricFileClient(download_url=None)  # type: ignore[arg-type]
    with pytest.raises(KricInvalidParameterError, match="data.kric.go.kr"):
        KricFileClient(download_url="https://127.0.0.1")
    with pytest.raises(KricInvalidParameterError, match="valid HTTPS URL"):
        KricFileClient(download_url="https://data.kric.go.kr:bad")
    for invalid_timeout in (None, True, "1", float("nan")):
        with pytest.raises(KricInvalidParameterError, match="timeout"):
            KricFileClient(timeout=invalid_timeout)  # type: ignore[arg-type]


def test_public_file_parser_normalizes_late_worksheet_xml_errors():
    source = _station_info_workbook_bytes()
    output = BytesIO()
    with ZipFile(BytesIO(source)) as source_archive, ZipFile(output, "w", ZIP_DEFLATED) as output_archive:
        for item in source_archive.infolist():
            data = source_archive.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = b"<worksheet>"
            output_archive.writestr(item, data)

    with pytest.raises(KricServerError, match="readable XLSX workbook"):
        parse_xlsx_table(output.getvalue())


@respx.mock
async def test_public_file_download_rejects_declared_oversize_before_reading_body():
    respx.get(FILE_DOWNLOAD_URL).mock(
        return_value=httpx.Response(200, content=b"small", headers={"content-length": "100"})
    )
    async with KricFileClient(max_download_bytes=10) as client:
        with pytest.raises(KricServerError, match="max_download_bytes"):
            await client.download_dataset(dataset_id=1294)
