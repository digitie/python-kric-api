"""KRIC 공개 파일 데이터의 무인증 다운로드와 XLSX 파싱을 제공한다."""

from __future__ import annotations

from collections.abc import Mapping
import hashlib
from io import BytesIO
from math import isfinite
import re
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

import httpx
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .exceptions import (
    KricInvalidParameterError,
    KricNetworkError,
    KricRateLimitError,
    KricServerError,
)
from .models import FileStationInfo, KricFileDownload, KricFileTable, StationCodeInfo
from .storage import RustfsObjectStore, StoredObject
from .parse import as_raw_mapping, float_or_none, require_fields, string_or_none

FILE_DOWNLOAD_URL = "https://data.kric.go.kr/rips/dataset/download.file"
KRIC_FILE_HOST = "data.kric.go.kr"
NATIONWIDE_STATION_INFO_DATASET_ID = 1294
STATION_CODE_NOTICE_ID = 17
STATION_CODE_FILE_ID = 1
STATION_CODE_FILE_URL = (
    "https://data.kric.go.kr/rips/download.file?type=L&id=17&answerId=17&fileId=1"
)
DEFAULT_MAX_DOWNLOAD_BYTES = 10 * 1024 * 1024
DEFAULT_MAX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
DEFAULT_MAX_ROWS = 100_000
DEFAULT_MAX_COLUMNS = 256

_STATION_INFO_REQUIRED_HEADERS = (
    "철도운영기관명",
    "운영노선",
    "역 번호",
    "역명(한글)",
)
_STATION_CODE_REQUIRED_HEADERS = (
    "RAIL_OPR_ISTT_CD",
    "RAIL_OPR_ISTT_NM",
    "LN_CD",
    "LN_NM",
    "STIN_CD",
    "STIN_NM",
)
_ZERO_PADDED_NUMBER_FORMAT = re.compile(r"(?P<prefix>(?:\\.)*)(?P<zeros>0+)$")


class KricFileClient:
    """KRIC 공개 파일 client. 인증키·Open API 요청을 사용하지 않는다."""

    def __init__(
        self,
        *,
        download_url: str = FILE_DOWNLOAD_URL,
        timeout: float = 30.0,
        max_download_bytes: int = DEFAULT_MAX_DOWNLOAD_BYTES,
        max_uncompressed_bytes: int = DEFAULT_MAX_UNCOMPRESSED_BYTES,
        max_rows: int = DEFAULT_MAX_ROWS,
        max_columns: int = DEFAULT_MAX_COLUMNS,
        client: httpx.AsyncClient | None = None,
    ) -> None:
        if (
            isinstance(timeout, bool)
            or not isinstance(timeout, (int, float))
            or not isfinite(timeout)
            or timeout <= 0
        ):
            raise KricInvalidParameterError("timeout must be positive")
        self.download_url = _https_url(download_url, "download_url")
        _positive_integer(max_download_bytes, "max_download_bytes")
        _positive_integer(max_uncompressed_bytes, "max_uncompressed_bytes")
        _positive_integer(max_rows, "max_rows")
        _positive_integer(max_columns, "max_columns")
        self.max_download_bytes = max_download_bytes
        self.max_uncompressed_bytes = max_uncompressed_bytes
        self.max_rows = max_rows
        self.max_columns = max_columns
        self._client = client or httpx.AsyncClient(timeout=timeout)
        self._owns_client = client is None

    async def __aenter__(self) -> "KricFileClient":
        return self

    async def __aexit__(self, *_: object) -> None:
        await self.aclose()

    async def aclose(self) -> None:
        """내부에서 만든 HTTP client만 닫는다."""
        if self._owns_client:
            await self._client.aclose()

    async def download_dataset(self, *, dataset_id: int, operation: int = 1) -> KricFileDownload:
        """공개 파일을 내려받는다. dataset/operation은 포털 상세 URL의 식별자다."""
        _positive_integer(dataset_id, "dataset_id")
        _positive_integer(operation, "operation")
        try:
            async with self._client.stream(
                "GET",
                self.download_url,
                params={"type": "filedata", "id": str(dataset_id), "operation": str(operation)},
                follow_redirects=False,
            ) as response:
                if response.status_code == 429:
                    raise KricRateLimitError("KRIC public-file request rate limited: HTTP 429")
                if 300 <= response.status_code < 400:
                    raise KricServerError(
                        f"KRIC public-file redirect denied: HTTP {response.status_code}"
                    )
                if response.status_code >= 400:
                    raise KricServerError(f"KRIC public-file request failed: HTTP {response.status_code}")
                _check_content_length(response.headers.get("content-length"), self.max_download_bytes)
                content = bytearray()
                async for chunk in response.aiter_bytes():
                    content.extend(chunk)
                    if len(content) > self.max_download_bytes:
                        raise KricServerError("KRIC public-file response exceeds max_download_bytes")
                source_url = str(response.request.url)
                content_type = response.headers.get("content-type")
                content_disposition = response.headers.get("content-disposition")
                etag = response.headers.get("etag")
        except httpx.HTTPError as exc:
            raise KricNetworkError("KRIC public-file request failed") from exc
        if not content:
            raise KricServerError("KRIC public-file response is empty")
        return KricFileDownload(
            dataset_id=dataset_id,
            operation=operation,
            source_url=source_url,
            content_type=content_type,
            content=bytes(content),
            content_disposition=content_disposition,
            etag=etag,
        )

    async def get_nationwide_station_info(self) -> tuple[FileStationInfo, ...]:
        """무인증 전국 도시광역철도 역사정보(XLSX, dataset 1294)를 파싱한다."""
        download = await self.download_dataset(dataset_id=NATIONWIDE_STATION_INFO_DATASET_ID)
        return parse_nationwide_station_info_xlsx(
            download.content,
            max_compressed_bytes=self.max_download_bytes,
            max_uncompressed_bytes=self.max_uncompressed_bytes,
            max_rows=self.max_rows,
            max_columns=self.max_columns,
        )

    async def get_station_codes(self) -> tuple[StationCodeInfo, ...]:
        """공식 자료실 역사 코드 XLSX를 인증키 없이 파싱한다."""
        download = await self._download_station_code_file()
        return parse_station_code_xlsx(
            download.content,
            max_compressed_bytes=self.max_download_bytes,
            max_uncompressed_bytes=self.max_uncompressed_bytes,
            max_rows=self.max_rows,
            max_columns=self.max_columns,
        )

    async def download_dataset_to_rustfs(
        self,
        store: RustfsObjectStore,
        *,
        dataset_id: int,
        operation: int = 1,
    ) -> tuple[KricFileDownload, StoredObject]:
        """공개 파일을 내려받아 형식을 단정하지 않고 공용 RustFS에 보관한다.

        object key는 dataset·operation·SHA-256에 의해 결정돼 같은 원문 재수집은
        같은 객체를 덮어쓴다. 범용 공개 파일은 XLSX라고 가정하지 않아 확장자 없는 key와
        provider가 보낸 `Content-Type`(없으면 `application/octet-stream`)을 사용한다.
        이 메서드는 파일 parse나 소비 서비스 DB 적재를 하지 않는다.
        """
        download = await self.download_dataset(dataset_id=dataset_id, operation=operation)
        return download, await self._archive_download(store, download)

    async def get_nationwide_station_info_to_rustfs(
        self, store: RustfsObjectStore
    ) -> tuple[tuple[FileStationInfo, ...], StoredObject]:
        """검증한 dataset 1294 XLSX만 RustFS에 보관하고 typed 역사정보를 반환한다."""
        download = await self.download_dataset(dataset_id=NATIONWIDE_STATION_INFO_DATASET_ID)
        stations = parse_nationwide_station_info_xlsx(
            download.content,
            max_compressed_bytes=self.max_download_bytes,
            max_uncompressed_bytes=self.max_uncompressed_bytes,
            max_rows=self.max_rows,
            max_columns=self.max_columns,
        )
        stored = await self._archive_download(
            store,
            download,
            suffix=".xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        return stations, stored

    async def get_station_codes_to_rustfs(
        self, store: RustfsObjectStore
    ) -> tuple[tuple[StationCodeInfo, ...], StoredObject]:
        """역사 코드 XLSX를 검증·파싱하고 checksum 기반 RustFS key에 원문을 보관한다."""
        download = await self._download_station_code_file()
        codes = parse_station_code_xlsx(
            download.content,
            max_compressed_bytes=self.max_download_bytes,
            max_uncompressed_bytes=self.max_uncompressed_bytes,
            max_rows=self.max_rows,
            max_columns=self.max_columns,
        )
        checksum = hashlib.sha256(download.content).hexdigest()
        stored = await store.put_bytes(
            object_key=store.prefixed_key(
                "kric", "station-codes", f"notice-{STATION_CODE_NOTICE_ID}",
                f"file-{STATION_CODE_FILE_ID}", f"{checksum}.xlsx",
            ),
            body=download.content,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        return codes, stored

    async def _download_station_code_file(self) -> KricFileDownload:
        """고정된 공식 자료실 첨부 파일만 받아 SSRF·redirect를 허용하지 않는다."""
        try:
            async with self._client.stream(
                "GET", STATION_CODE_FILE_URL, follow_redirects=False
            ) as response:
                if response.status_code == 429:
                    raise KricRateLimitError("KRIC station-code file request rate limited: HTTP 429")
                if 300 <= response.status_code < 400:
                    raise KricServerError(
                        f"KRIC station-code file redirect denied: HTTP {response.status_code}"
                    )
                if response.status_code >= 400:
                    raise KricServerError(
                        f"KRIC station-code file request failed: HTTP {response.status_code}"
                    )
                _check_content_length(response.headers.get("content-length"), self.max_download_bytes)
                content = bytearray()
                async for chunk in response.aiter_bytes():
                    content.extend(chunk)
                    if len(content) > self.max_download_bytes:
                        raise KricServerError("KRIC station-code file response exceeds max_download_bytes")
                source_url = str(response.request.url)
                content_type = response.headers.get("content-type")
                content_disposition = response.headers.get("content-disposition")
                etag = response.headers.get("etag")
        except httpx.HTTPError as exc:
            raise KricNetworkError("KRIC station-code file request failed") from exc
        if not content:
            raise KricServerError("KRIC station-code file response is empty")
        return KricFileDownload(
            dataset_id=STATION_CODE_NOTICE_ID,
            operation=STATION_CODE_FILE_ID,
            source_url=source_url,
            content_type=content_type,
            content=bytes(content),
            content_disposition=content_disposition,
            etag=etag,
        )

    async def _archive_download(
        self,
        store: RustfsObjectStore,
        download: KricFileDownload,
        *,
        suffix: str = "",
        content_type: str | None = None,
    ) -> StoredObject:
        checksum = hashlib.sha256(download.content).hexdigest()
        return await store.put_bytes(
            object_key=store.prefixed_key(
                "kric",
                f"dataset-{download.dataset_id}",
                f"operation-{download.operation}",
                f"{checksum}{suffix}",
            ),
            body=download.content,
            content_type=content_type or download.content_type or "application/octet-stream",
        )


def parse_nationwide_station_info_xlsx(
    content: bytes,
    *,
    max_compressed_bytes: int = DEFAULT_MAX_DOWNLOAD_BYTES,
    max_uncompressed_bytes: int = DEFAULT_MAX_UNCOMPRESSED_BYTES,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_columns: int = DEFAULT_MAX_COLUMNS,
) -> tuple[FileStationInfo, ...]:
    """dataset 1294의 첫 worksheet를 typed 역사 기준정보로 변환한다."""
    table = parse_xlsx_table(
        content,
        max_compressed_bytes=max_compressed_bytes,
        max_uncompressed_bytes=max_uncompressed_bytes,
        max_rows=max_rows,
        max_columns=max_columns,
    )
    missing = [header for header in _STATION_INFO_REQUIRED_HEADERS if header not in table.headers]
    if missing:
        raise KricServerError(
            "KRIC station-info workbook is missing required headers: " + ", ".join(missing)
        )
    return tuple(parse_nationwide_station_info_row(row) for row in table.rows)


def parse_station_code_xlsx(
    content: bytes,
    *,
    max_compressed_bytes: int = DEFAULT_MAX_DOWNLOAD_BYTES,
    max_uncompressed_bytes: int = DEFAULT_MAX_UNCOMPRESSED_BYTES,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_columns: int = DEFAULT_MAX_COLUMNS,
) -> tuple[StationCodeInfo, ...]:
    """자료실의 첫 worksheet 역사 코드 행을 API 요청용 typed 식별자로 변환한다."""
    table = parse_xlsx_table(
        content,
        max_compressed_bytes=max_compressed_bytes,
        max_uncompressed_bytes=max_uncompressed_bytes,
        max_rows=max_rows,
        max_columns=max_columns,
    )
    missing = [header for header in _STATION_CODE_REQUIRED_HEADERS if header not in table.headers]
    if missing:
        raise KricServerError(
            "KRIC station-code workbook is missing required headers: " + ", ".join(missing)
        )
    rows = tuple(parse_station_code_row(row) for row in table.rows)
    if not rows:
        raise KricServerError("KRIC station-code workbook contains no code rows")
    return rows


def parse_xlsx_table(
    content: bytes,
    *,
    max_compressed_bytes: int = DEFAULT_MAX_DOWNLOAD_BYTES,
    max_uncompressed_bytes: int = DEFAULT_MAX_UNCOMPRESSED_BYTES,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_columns: int = DEFAULT_MAX_COLUMNS,
) -> KricFileTable:
    """공개 XLSX의 첫 worksheet를 헤더 기반 원문 table로 읽는다.

    시설별 파일처럼 아직 typed 계약이 확정되지 않은 자료도 이 함수로 안전하게 소비할 수 있다.
    """
    if not isinstance(content, bytes):
        raise KricInvalidParameterError("content must be bytes")
    _positive_integer(max_compressed_bytes, "max_compressed_bytes")
    _positive_integer(max_uncompressed_bytes, "max_uncompressed_bytes")
    _positive_integer(max_rows, "max_rows")
    _positive_integer(max_columns, "max_columns")
    if len(content) > max_compressed_bytes:
        raise KricServerError("KRIC public-file response exceeds max_compressed_bytes")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(member.file_size for member in archive.infolist()) > max_uncompressed_bytes:
                raise KricServerError("KRIC public-file workbook exceeds max_uncompressed_bytes")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ParseError, ValueError) as exc:
        raise KricServerError("KRIC public file is not a readable XLSX workbook") from exc
    try:
        if not workbook.worksheets:
            raise KricServerError("KRIC public-file workbook has no worksheet")
        sheet = workbook.worksheets[0]
        if sheet.max_row > max_rows:
            raise KricServerError("KRIC public-file workbook exceeds max_rows")
        if sheet.max_column > max_columns:
            raise KricServerError("KRIC public-file workbook exceeds max_columns")
        header_rows = sheet.iter_rows(values_only=True, max_row=1)
        try:
            header_values = next(header_rows)
        except StopIteration as exc:
            raise KricServerError("KRIC public-file workbook is empty") from exc
        headers = tuple(_header_text(value) for value in header_values)
        if not headers or any(header is None for header in headers):
            raise KricServerError("KRIC public-file workbook has a blank header")
        normalized_headers = tuple(header for header in headers if header is not None)
        if len(set(normalized_headers)) != len(normalized_headers):
            raise KricServerError("KRIC public-file workbook has duplicate headers")
        records: list[Mapping[str, str | None]] = []
        for row_index, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
            if row_index > max_rows:
                raise KricServerError("KRIC public-file workbook exceeds max_rows")
            raw_row = {
                header: _cell_value(cell)
                for header, cell in zip(normalized_headers, cells, strict=True)
            }
            if all(value is None or str(value).strip() == "" for value in raw_row.values()):
                continue
            records.append(as_raw_mapping(raw_row))
        return KricFileTable(
            worksheet_title=sheet.title,
            headers=normalized_headers,
            rows=tuple(records),
        )
    except ParseError as exc:
        raise KricServerError("KRIC public-file workbook contains invalid worksheet XML") from exc
    finally:
        workbook.close()


def parse_nationwide_station_info_row(row: Mapping[str, Any]) -> FileStationInfo:
    """dataset 1294의 한 행을 파싱한다. 명세 밖 열도 raw에 보존한다."""
    raw = as_raw_mapping(row)
    require_fields(raw, "station-info file row", *_STATION_INFO_REQUIRED_HEADERS)
    return FileStationInfo(
        rail_operator_name=raw.get("철도운영기관명"),
        operating_line_name=raw.get("운영노선"),
        station_type=raw.get("역 종류"),
        station_number=raw.get("역 번호"),
        station_name=raw.get("역명(한글)"),
        english_name=raw.get("역명(영어)"),
        romanized_name=raw.get("역명(로마자)"),
        japanese_name=raw.get("역명(일본어)"),
        simplified_chinese_name=raw.get("역명(중국어간체)"),
        traditional_chinese_name=raw.get("역명(중국어번체)"),
        sub_station_name=raw.get("역명(부역명)"),
        longitude=float_or_none(row.get("역 위치(경도)"), "역 위치(경도)"),
        latitude=float_or_none(row.get("역 위치(위도)"), "역 위치(위도)"),
        lot_address=raw.get("역 주소(지번주소)"),
        road_address=raw.get("역 주소(도로명 주소)"),
        station_phone_number=raw.get("역사 전화번호"),
        data_reference_date=raw.get("데이터 기준일자"),
        raw=raw,
    )


def parse_station_code_row(row: Mapping[str, Any]) -> StationCodeInfo:
    """역사 코드 파일 한 행의 원시 코드와 표시명을 함께 보존한다."""
    raw = as_raw_mapping(row)
    require_fields(raw, "station-code file row", *_STATION_CODE_REQUIRED_HEADERS)
    return StationCodeInfo(
        rail_operator_code=raw.get("RAIL_OPR_ISTT_CD"),
        rail_operator_name=raw.get("RAIL_OPR_ISTT_NM"),
        line_code=raw.get("LN_CD"),
        line_name=raw.get("LN_NM"),
        station_code=raw.get("STIN_CD"),
        station_name=raw.get("STIN_NM"),
        raw=raw,
    )


def _header_text(value: object) -> str | None:
    return string_or_none(value)


def _cell_value(cell: Any) -> object:
    """0 패딩과 이스케이프된 접두사가 있는 수치 서식을 표시 문자열로 보존한다."""
    value = cell.value
    number_format = cell.number_format
    zero_padded = (
        _ZERO_PADDED_NUMBER_FORMAT.fullmatch(number_format)
        if isinstance(number_format, str)
        else None
    )
    if (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and float(value).is_integer()
        and zero_padded is not None
    ):
        prefix = re.sub(r"\\(.)", r"\1", zero_padded.group("prefix"))
        return f"{prefix}{int(value):0{len(zero_padded.group('zeros'))}d}"
    return value


def _https_url(value: object, name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise KricInvalidParameterError(f"{name} must be a non-blank HTTPS URL")
    try:
        parsed = httpx.URL(value)
    except (TypeError, httpx.InvalidURL) as exc:
        raise KricInvalidParameterError(f"{name} must be a valid HTTPS URL") from exc
    if parsed.scheme != "https" or parsed.host != KRIC_FILE_HOST or parsed.port not in (None, 443):
        raise KricInvalidParameterError(f"{name} must be an HTTPS URL for {KRIC_FILE_HOST}")
    return str(parsed)


def _check_content_length(value: str | None, maximum: int) -> None:
    if value is None:
        return
    try:
        length = int(value)
    except ValueError as exc:
        raise KricServerError("KRIC public-file response has invalid Content-Length") from exc
    if length < 0:
        raise KricServerError("KRIC public-file response has invalid Content-Length")
    if length > maximum:
        raise KricServerError("KRIC public-file response exceeds max_download_bytes")


def _positive_integer(value: int, name: str) -> None:
    if not isinstance(value, int) or isinstance(value, bool) or value <= 0:
        raise KricInvalidParameterError(f"{name} must be a positive integer")
